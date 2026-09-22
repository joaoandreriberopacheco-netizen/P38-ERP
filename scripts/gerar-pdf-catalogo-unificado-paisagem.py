#!/usr/bin/env python3
"""PDF catálogo unificado — A4 paisagem, todas as camadas + SKU Supabase (nome)."""

from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
PAGE_SIZE = landscape(A4)
LINE = colors.HexColor("#E5E5E5")
TEXT = colors.HexColor("#1F1F1F")
MUTED = colors.HexColor("#6B6B6B")
HEADER_BG = colors.HexColor("#FAFAFA")

COL_HEADERS = [
    "codigo_4x",
    "etapa",
    "categoria",
    "subcategoria",
    "linha",
    "comp1",
    "comp2",
    "comp3",
    "sku",
    "sku_supabase",
]
COL_WIDTHS = [
    12 * mm,
    18 * mm,
    20 * mm,
    18 * mm,
    22 * mm,
    34 * mm,
    28 * mm,
    18 * mm,
    14 * mm,
    52 * mm,
]
MAX_ROWS = 22


def cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def load_filter(filter_path: Path) -> tuple[set[str], dict[str, str], str, str]:
    data = json.loads(filter_path.read_text(encoding="utf-8"))
    codigos = {str(c).strip().upper() for c in data.get("codigos", []) if str(c).strip()}
    nomes = {
        str(k).strip().upper(): cell_str(v)
        for k, v in (data.get("nomesSupabase") or {}).items()
    }
    kind = str(data.get("kind", "filtro")).strip()
    cutoff = str(data.get("cutoff", "")).strip()
    return codigos, nomes, kind, cutoff


def load_rows(xlsx_path: Path, codigos: set[str], nomes: dict[str, str]) -> list[dict[str, str]]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["Catálogo 4×3"]
    rows: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[9]:
            continue
        sku = cell_str(row[9]).upper()
        if sku not in codigos:
            continue
        rows.append(
            {
                "codigo_4x": cell_str(row[0]),
                "etapa": cell_str(row[2]),
                "categoria": cell_str(row[3]),
                "subcategoria": cell_str(row[4]),
                "linha": cell_str(row[5]),
                "comp1": cell_str(row[6]),
                "comp2": cell_str(row[7]),
                "comp3": cell_str(row[8]),
                "sku": sku,
                "sku_supabase": nomes.get(sku) or "—",
            }
        )
    wb.close()
    rows.sort(key=lambda r: (r["etapa"], r["categoria"], r["subcategoria"], r["linha"], r["comp1"], r["sku"]))
    return rows


def build_pdf(rows: list[dict[str, str]], out_path: Path, *, kind: str, cutoff: str) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=PAGE_SIZE,
        leftMargin=4 * mm,
        rightMargin=4 * mm,
        topMargin=5 * mm,
        bottomMargin=6 * mm,
        title="P38 — Catálogo unificado",
        author="P38 ERP",
    )
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=6.5, leading=8, textColor=TEXT)
    code = ParagraphStyle("code", parent=body, fontName="Courier", fontSize=6, textColor=MUTED)
    header = ParagraphStyle("header", fontName="Helvetica-Bold", fontSize=6.5, leading=8, textColor=TEXT)
    title = ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=TEXT, spaceAfter=2)
    subtitle = ParagraphStyle("subtitle", fontName="Helvetica", fontSize=8, leading=10, textColor=MUTED, spaceAfter=4)
    generated = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
    story = [
        Paragraph("P38 — Catálogo unificado (paisagem)", title),
        Paragraph(
            f"A4 paisagem · {len(rows)} SKUs · {kind} · desde {cutoff or '—'} · {generated}",
            subtitle,
        ),
    ]

    def cell(text: str, style: ParagraphStyle):
        safe = cell_str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;") or "—"
        return Paragraph(safe, style)

    for i in range(0, len(rows), MAX_ROWS):
        chunk = rows[i : i + MAX_ROWS]
        table_data = [[Paragraph(h, header) for h in COL_HEADERS]]
        for r in chunk:
            table_data.append(
                [
                    cell(r["codigo_4x"], code),
                    cell(r["etapa"], body),
                    cell(r["categoria"], body),
                    cell(r["subcategoria"], body),
                    cell(r["linha"], body),
                    cell(r["comp1"], body),
                    cell(r["comp2"], body),
                    cell(r["comp3"], body),
                    cell(r["sku"], code),
                    cell(r["sku_supabase"], body),
                ]
            )
        table = Table(table_data, colWidths=COL_WIDTHS, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), HEADER_BG),
                    ("GRID", (0, 0), (-1, -1), 0.25, LINE),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 2),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )
        story.append(table)
        if i + MAX_ROWS < len(rows):
            story.append(Spacer(1, 3 * mm))

    def footer(canvas, doc_obj):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(MUTED)
        canvas.drawString(4 * mm, 4 * mm, f"P38 · unificado paisagem · {len(rows)} SKUs")
        canvas.drawRightString(PAGE_SIZE[0] - 4 * mm, 4 * mm, f"pág. {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def main() -> int:
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs/exports/P38-catalogo-4x3.xlsx"
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "docs/exports/P38-catalogo-unificado-paisagem.pdf"
    filter_path = Path(sys.argv[3]) if len(sys.argv) > 3 else None
    if not filter_path or not filter_path.is_file():
        print("Filtro JSON em falta", file=sys.stderr)
        return 1
    codigos, nomes, kind, cutoff = load_filter(filter_path)
    rows = load_rows(xlsx, codigos, nomes)
    build_pdf(rows, out, kind=kind, cutoff=cutoff)
    print(f"[pdf:unificado-paisagem] {len(rows)} SKUs → {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
