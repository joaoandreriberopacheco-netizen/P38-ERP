#!/usr/bin/env python3
"""PDF catálogo 4× — nível drill (etapa → produto compra + qtd SKUs). Estilo Cursor."""

from __future__ import annotations

import json
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "docs" / "exports" / "P38-catalogo-4x3.xlsx"
DEFAULT_OUT = ROOT / "docs" / "exports" / "P38-catalogo-4x-nivel.pdf"
ARTIFACT_OUT = Path("/opt/cursor/artifacts/P38-catalogo-4x-nivel.pdf")

LINE = colors.HexColor("#E5E5E5")
TEXT = colors.HexColor("#1F1F1F")
MUTED = colors.HexColor("#6B6B6B")
HEADER_BG = colors.HexColor("#FAFAFA")

COL_HEADERS = [
    "etapa",
    "categoria",
    "subcategoria",
    "linha",
    "produto compra",
    "skus",
]
COL_WIDTHS = [22 * mm, 28 * mm, 26 * mm, 34 * mm, 64 * mm, 14 * mm]
MERGE_KEYS = ("etapa", "categoria", "subcategoria", "linha")
MERGE_COLS = tuple(COL_HEADERS.index(k) for k in MERGE_KEYS)
MAX_ROWS_PER_TABLE = 24
PAGE_MARGIN_X = 4 * mm
PAGE_MARGIN_TOP = 5 * mm
PAGE_MARGIN_BOTTOM = 6 * mm
PAGE_SIZE = A4


def cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def produto_key_tuple(
    etapa: str, categoria: str, subcategoria: str, linha: str, produto: str
) -> str:
    return "\x00".join(
        cell_str(v)
        for v in (etapa, categoria, subcategoria, linha, produto)
    )


def load_filter_produto_keys(filter_path: Path | None) -> set[str] | None:
    if not filter_path or not filter_path.is_file():
        return None
    data = json.loads(filter_path.read_text(encoding="utf-8"))
    return {str(k).strip() for k in data.get("produto_keys", []) if str(k).strip()}


def load_filter_kind(filter_path: Path | None) -> str | None:
    if not filter_path or not filter_path.is_file():
        return None
    data = json.loads(filter_path.read_text(encoding="utf-8"))
    return str(data.get("kind", "sem-estoque")).strip() or "sem-estoque"


def filter_copy(kind: str | None) -> tuple[str, str, str]:
    if kind == "zumbis":
        return (
            "P38 — Catálogo 4× (zumbis)",
            "A4 retrato · {n} produtos compra zumbis · {skus} SKUs · sem mov. 4 meses · {generated}",
            " · zumbis",
        )
    if kind == "sem-venda-75d":
        return (
            "P38 — Catálogo 4× (sem estoque · sem venda 75d)",
            "A4 retrato · {n} produtos compra · {skus} SKUs · sem venda 75 dias · {generated}",
            " · sem venda 75d",
        )
    if kind == "sem-movimento-45d":
        return (
            "P38 — Catálogo 4× (sem estoque · sem mov. 45d)",
            "A4 retrato · {n} produtos compra · {skus} SKUs · sem mov. 45 dias · {generated}",
            " · sem mov. 45d",
        )
    return (
        "P38 — Catálogo 4× (sem estoque)",
        "A4 retrato · {n} produtos compra sem estoque · {skus} SKUs · {generated}",
        " · sem estoque",
    )


def load_and_aggregate(
    xlsx_path: Path, *, produto_keys_filter: set[str] | None = None
) -> list[dict[str, str]]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["Catálogo 4×3"]
    counts: dict[tuple[str, str, str, str, str], int] = defaultdict(int)

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[9]:
            continue
        produto = cell_str(row[6]) or cell_str(row[12]) or cell_str(row[7]) or "(sem produto compra)"
        etapa = cell_str(row[2])
        categoria = cell_str(row[3])
        subcategoria = cell_str(row[4])
        linha = cell_str(row[5])
        pk = produto_key_tuple(etapa, categoria, subcategoria, linha, produto)
        if produto_keys_filter is not None and pk not in produto_keys_filter:
            continue
        key = (etapa, categoria, subcategoria, linha, produto)
        counts[key] += 1
    wb.close()

    aggregated = [
        {
            "etapa": k[0],
            "categoria": k[1],
            "subcategoria": k[2],
            "linha": k[3],
            "produto_compra": k[4],
            "skus": str(n),
        }
        for k, n in counts.items()
    ]
    aggregated.sort(
        key=lambda r: (
            r["etapa"],
            r["categoria"],
            r["subcategoria"],
            r["linha"],
            r["produto_compra"],
        )
    )
    return aggregated


def collapse_parents(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    out: list[dict[str, str]] = []
    prev = {k: "" for k in MERGE_KEYS}
    for row in rows:
        collapsed = dict(row)
        for key in MERGE_KEYS:
            if collapsed[key] == prev[key]:
                collapsed[key] = ""
            else:
                prev[key] = collapsed[key]
        out.append(collapsed)
    return out


def fill_parent_context(row: dict[str, str], context: dict[str, str]) -> dict[str, str]:
    filled = dict(row)
    for key in MERGE_KEYS:
        if not filled[key] and context.get(key):
            filled[key] = context[key]
    return filled


def chunk_group_rows(group_rows: list[dict[str, str]]) -> list[list[dict[str, str]]]:
    if len(group_rows) <= MAX_ROWS_PER_TABLE:
        return [group_rows]

    chunks: list[list[dict[str, str]]] = []
    context = {k: "" for k in MERGE_KEYS}
    for i in range(0, len(group_rows), MAX_ROWS_PER_TABLE):
        slice_rows = group_rows[i : i + MAX_ROWS_PER_TABLE]
        chunk: list[dict[str, str]] = []
        for j, row in enumerate(slice_rows):
            chunk.append(fill_parent_context(row, context) if j == 0 else dict(row))
        chunks.append(chunk)
        for key in MERGE_KEYS:
            for row in reversed(group_rows[: i + len(slice_rows)]):
                if row[key]:
                    context[key] = row[key]
                    break
    return chunks


def group_by_etapa(rows: list[dict[str, str]]) -> list[list[dict[str, str]]]:
    groups: list[list[dict[str, str]]] = []
    current: list[dict[str, str]] = []
    for row in rows:
        if row["etapa"] and current:
            groups.append(current)
            current = []
        current.append(row)
    if current:
        groups.append(current)
    return groups


def compute_vertical_spans(
    rows: list[dict[str, str]], key: str, col: int, *, row_offset: int
) -> list[tuple]:
    spans: list[tuple] = []
    i = 0
    while i < len(rows):
        if not rows[i][key]:
            i += 1
            continue
        start = i
        j = i + 1
        while j < len(rows) and not rows[j][key]:
            j += 1
        if j - start > 1:
            spans.append(("SPAN", (col, start + row_offset), (col, (j - 1) + row_offset)))
        i = j
    return spans


def build_pdf(
    rows: list[dict[str, str]],
    out_path: Path,
    total_skus: int,
    *,
    sem_estoque: bool = False,
    filter_kind: str | None = None,
) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=PAGE_SIZE,
        leftMargin=PAGE_MARGIN_X,
        rightMargin=PAGE_MARGIN_X,
        topMargin=PAGE_MARGIN_TOP,
        bottomMargin=PAGE_MARGIN_BOTTOM,
        title="P38 — Catálogo 4× (nível drill)",
        author="P38 ERP",
    )

    body = ParagraphStyle(
        "body",
        fontName="Helvetica",
        fontSize=7.5,
        leading=9,
        textColor=TEXT,
        alignment=TA_LEFT,
    )
    count_style = ParagraphStyle(
        "count",
        parent=body,
        fontName="Courier",
        fontSize=7,
        textColor=TEXT,
        alignment=TA_RIGHT,
    )
    header = ParagraphStyle(
        "header",
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9,
        textColor=TEXT,
        alignment=TA_LEFT,
    )
    header_count = ParagraphStyle(
        "header_count",
        parent=header,
        alignment=TA_RIGHT,
    )
    title = ParagraphStyle(
        "title",
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=TEXT,
        spaceAfter=2,
    )
    subtitle = ParagraphStyle(
        "subtitle",
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=MUTED,
        spaceAfter=4,
    )

    generated = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
    if sem_estoque:
        titulo, subt_tpl, footer_tag = filter_copy(filter_kind)
        subtitulo = subt_tpl.format(n=len(rows), skus=total_skus, generated=generated)
    else:
        titulo = "P38 — Catálogo 4× (nível drill)"
        subtitulo = f"A4 retrato · {len(rows)} produtos compra · {total_skus} SKUs · {generated}"
        footer_tag = ""
    story = [
        Paragraph(titulo, title),
        Paragraph(subtitulo, subtitle),
    ]

    def cell(text: str, style: ParagraphStyle, *, blank_if_empty: bool = False):
        raw = text.replace("\n", " ").strip()
        if not raw:
            if blank_if_empty:
                return ""
            raw = "—"
        safe = raw.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(safe, style)

    def build_group_table(group_rows: list[dict[str, str]], *, include_header: bool) -> Table:
        table_data: list[list] = []
        if include_header:
            table_data.append(
                [
                    cell(h, header_count if h == "skus" else header)
                    for h in COL_HEADERS
                ]
            )
        for row in group_rows:
            table_data.append(
                [
                    cell(row["etapa"], body, blank_if_empty=True),
                    cell(row["categoria"], body, blank_if_empty=True),
                    cell(row["subcategoria"], body, blank_if_empty=True),
                    cell(row["linha"], body, blank_if_empty=True),
                    cell(row["produto_compra"], body),
                    cell(row["skus"], count_style),
                ]
            )

        row_offset = 1 if include_header else 0
        span_styles: list[tuple] = []
        for key, col in zip(MERGE_KEYS, MERGE_COLS):
            span_styles.extend(
                compute_vertical_spans(group_rows, key, col, row_offset=row_offset)
            )

        repeat_rows = 1 if include_header else 0
        table = Table(table_data, colWidths=COL_WIDTHS, repeatRows=repeat_rows)
        data_start = 1 if include_header else 0
        table.setStyle(
            TableStyle(
                [
                    *(
                        [("BACKGROUND", (0, 0), (-1, 0), HEADER_BG)]
                        if include_header
                        else []
                    ),
                    ("VALIGN", (0, data_start), (-1, -1), "TOP"),
                    ("VALIGN", (0, data_start), (3, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.25, LINE),
                    ("LINEAFTER", (0, 0), (-2, -1), 0.25, LINE),
                    *span_styles,
                ]
            )
        )
        return table

    groups = group_by_etapa(rows)
    header_used = False
    for group_rows in groups:
        for chunk in chunk_group_rows(group_rows):
            story.append(build_group_table(chunk, include_header=not header_used))
            header_used = True
            story.append(Spacer(1, 1 * mm))

    def footer(canvas, doc_obj):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(MUTED)
        canvas.drawString(
            doc_obj.leftMargin,
            4 * mm,
            f"P38 · Catálogo 4×{footer_tag if sem_estoque else ' drill'} · {len(rows)} produtos · {total_skus} SKUs",
        )
        canvas.drawRightString(
            PAGE_SIZE[0] - doc_obj.rightMargin,
            4 * mm,
            f"pág. {canvas.getPageNumber()}",
        )
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def main() -> int:
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUT
    filter_path = Path(sys.argv[3]) if len(sys.argv) > 3 else None

    if not xlsx.is_file():
        print(f"Ficheiro em falta: {xlsx}", file=sys.stderr)
        return 1

    produto_keys_filter = load_filter_produto_keys(filter_path)
    filter_kind = load_filter_kind(filter_path)
    sem_estoque = produto_keys_filter is not None
    aggregated = load_and_aggregate(xlsx, produto_keys_filter=produto_keys_filter)
    rows = collapse_parents(aggregated)
    total_skus = sum(int(r["skus"]) for r in aggregated)
    build_pdf(rows, out, total_skus, sem_estoque=sem_estoque, filter_kind=filter_kind)
    label = filter_kind or ("sem-estoque-nivel" if sem_estoque else "catalogo-4x-nivel")
    print(f"[pdf:{label}] {len(rows)} produtos compra · {total_skus} SKUs → {out}")

    if not sem_estoque:
        try:
            ARTIFACT_OUT.parent.mkdir(parents=True, exist_ok=True)
            ARTIFACT_OUT.write_bytes(out.read_bytes())
            print(f"[pdf:catalogo-4x-nivel] cópia → {ARTIFACT_OUT}")
        except OSError:
            pass

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
