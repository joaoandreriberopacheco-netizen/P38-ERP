#!/usr/bin/env python3
"""PDF catálogo 4×3 — estilo tabela Cursor (modo claro, linhas finas, sem tampas laterais)."""

from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "docs" / "exports" / "P38-catalogo-4x3.xlsx"
DEFAULT_OUT = ROOT / "docs" / "exports" / "P38-catalogo-4x3.pdf"
ARTIFACT_OUT = Path("/opt/cursor/artifacts/P38-catalogo-4x3.pdf")

LINE = colors.HexColor("#E5E5E5")
TEXT = colors.HexColor("#1F1F1F")
MUTED = colors.HexColor("#6B6B6B")
HEADER_BG = colors.HexColor("#FAFAFA")

COL_HEADERS = ["codigo_4x", "linha", "comp1", "comp2", "comp3", "sku"]
# A4 retrato 210 mm — margens 4 mm → ~202 mm úteis
COL_WIDTHS = [16 * mm, 32 * mm, 56 * mm, 48 * mm, 26 * mm, 24 * mm]
MERGE_KEYS = ("codigo_4x", "linha", "comp1")
MERGE_COLS = tuple(COL_HEADERS.index(k) for k in MERGE_KEYS)
MAX_ROWS_PER_TABLE = 20
PAGE_MARGIN_X = 4 * mm
PAGE_MARGIN_TOP = 5 * mm
PAGE_MARGIN_BOTTOM = 6 * mm
PAGE_SIZE = A4  # retrato explícito (595×842 pt)


def cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def load_filter_codigos(filter_path: Path | None) -> set[str] | None:
    if not filter_path or not filter_path.is_file():
        return None
    data = json.loads(filter_path.read_text(encoding="utf-8"))
    return {str(c).strip().upper() for c in data.get("codigos", []) if str(c).strip()}


def load_filter_kind(filter_path: Path | None) -> str | None:
    if not filter_path or not filter_path.is_file():
        return None
    data = json.loads(filter_path.read_text(encoding="utf-8"))
    return str(data.get("kind", "sem-estoque")).strip() or "sem-estoque"


def filter_copy(kind: str | None) -> tuple[str, str, str]:
    if kind == "zumbis":
        return (
            "P38 — Catálogo 4×3 (zumbis)",
            "A4 retrato · {n} SKUs · zumbis · sem estoque · sem mov. 4 meses · {generated}",
            " · zumbis",
        )
    return (
        "P38 — Catálogo 4×3 (sem estoque)",
        "A4 retrato · {n} SKUs · produtos compra sem estoque · {generated}",
        " · sem estoque",
    )


def load_catalog_rows(
    xlsx_path: Path, *, codigos_filter: set[str] | None = None
) -> list[dict[str, str]]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["Catálogo 4×3"]
    rows: list[dict[str, str]] = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if not row or not row[9]:
            continue
        codigo = cell_str(row[9]).upper()
        if codigos_filter is not None and codigo not in codigos_filter:
            continue
        rows.append(
            {
                "codigo_4x": cell_str(row[0]),
                "linha": cell_str(row[5]),
                "comp1": cell_str(row[6]),
                "comp2": cell_str(row[7]),
                "comp3": cell_str(row[8]),
                "sku": codigo,
            }
        )
    wb.close()
    return rows


def collapse_groups(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    out: list[dict[str, str]] = []
    prev = {"codigo_4x": "", "linha": "", "comp1": ""}
    for row in rows:
        collapsed = dict(row)
        for key in ("codigo_4x", "linha", "comp1"):
            if collapsed[key] == prev[key]:
                collapsed[key] = ""
            else:
                prev[key] = collapsed[key]
        out.append(collapsed)
    return out


def fill_parent_context(row: dict[str, str], context: dict[str, str]) -> dict[str, str]:
    filled = dict(row)
    for key in MERGE_KEYS:
        if not filled[key] and context.get(key):
            filled[key] = context[key]
    return filled


def chunk_group_rows(group_rows: list[dict[str, str]]) -> list[list[dict[str, str]]]:
    """Parte blocos grandes para caber na página; repõe células-mãe no início de cada parte."""
    if len(group_rows) <= MAX_ROWS_PER_TABLE:
        return [group_rows]

    chunks: list[list[dict[str, str]]] = []
    context = {"codigo_4x": "", "linha": "", "comp1": ""}
    for i in range(0, len(group_rows), MAX_ROWS_PER_TABLE):
        slice_rows = group_rows[i : i + MAX_ROWS_PER_TABLE]
        chunk: list[dict[str, str]] = []
        for j, row in enumerate(slice_rows):
            if j == 0:
                chunk.append(fill_parent_context(row, context))
            else:
                chunk.append(dict(row))
        chunks.append(chunk)
        for key in MERGE_KEYS:
            for row in reversed(group_rows[: i + len(slice_rows)]):
                if row[key]:
                    context[key] = row[key]
                    break
    return chunks


def group_by_codigo(rows: list[dict[str, str]]) -> list[list[dict[str, str]]]:
    groups: list[list[dict[str, str]]] = []
    current: list[dict[str, str]] = []
    for row in rows:
        if row["codigo_4x"] and current:
            groups.append(current)
            current = []
        current.append(row)
    if current:
        groups.append(current)
    return groups


def compute_vertical_spans(
    rows: list[dict[str, str]], key: str, col: int, *, row_offset: int
) -> list[tuple]:
    """Rowspan para células-mãe (filhas ficam vazias e entram no SPAN)."""
    spans: list[tuple] = []
    i = 0
    while i < len(rows):
        if not rows[i][key]:
            i += 1
            continue
        start = i
        j = i + 1
        while j < len(rows) and not rows[j][key]:
            j += 1
        if j - start > 1:
            start_row = start + row_offset
            end_row = (j - 1) + row_offset
            spans.append(("SPAN", (col, start_row), (col, end_row)))
        i = j
    return spans


def build_pdf(
    rows: list[dict[str, str]],
    out_path: Path,
    *,
    sem_estoque: bool = False,
    filter_kind: str | None = None,
) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=PAGE_SIZE,
        leftMargin=PAGE_MARGIN_X,
        rightMargin=PAGE_MARGIN_X,
        topMargin=PAGE_MARGIN_TOP,
        bottomMargin=PAGE_MARGIN_BOTTOM,
        title="P38 — Catálogo 4×3",
        author="P38 ERP",
    )

    body = ParagraphStyle(
        "body",
        fontName="Helvetica",
        fontSize=7.5,
        leading=9,
        textColor=TEXT,
        alignment=TA_LEFT,
    )
    code = ParagraphStyle(
        "code",
        parent=body,
        fontName="Courier",
        fontSize=7,
        textColor=MUTED,
    )
    header = ParagraphStyle(
        "header",
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9,
        textColor=TEXT,
        alignment=TA_LEFT,
    )
    title = ParagraphStyle(
        "title",
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=TEXT,
        spaceAfter=2,
    )
    subtitle = ParagraphStyle(
        "subtitle",
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=MUTED,
        spaceAfter=4,
    )

    generated = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
    if sem_estoque:
        titulo, subt_tpl, footer_tag = filter_copy(filter_kind)
        subtitulo = subt_tpl.format(n=len(rows), generated=generated)
    else:
        titulo = "P38 — Catálogo 4×3"
        subtitulo = f"A4 retrato · {len(rows)} SKUs · coluna final = código · {generated}"
        footer_tag = ""
    story = [
        Paragraph(titulo, title),
        Paragraph(subtitulo, subtitle),
    ]

    def cell(text: str, style: ParagraphStyle, *, blank_if_empty: bool = False):
        raw = text.replace("\n", " ").strip()
        if not raw:
            if blank_if_empty:
                return ""
            raw = "—"
        safe = raw.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(safe, style)

    def build_group_table(group_rows: list[dict[str, str]], *, include_header: bool) -> Table:
        table_data: list[list] = []
        if include_header:
            table_data.append([cell(h, header) for h in COL_HEADERS])
        for row in group_rows:
            table_data.append(
                [
                    cell(row["codigo_4x"], code, blank_if_empty=True),
                    cell(row["linha"], body, blank_if_empty=True),
                    cell(row["comp1"], body, blank_if_empty=True),
                    cell(row["comp2"], body),
                    cell(row["comp3"], body if row["comp3"] else code),
                    cell(row["sku"], code),
                ]
            )

        row_offset = 1 if include_header else 0
        span_styles: list[tuple] = []
        for key, col in zip(MERGE_KEYS, MERGE_COLS):
            span_styles.extend(
                compute_vertical_spans(group_rows, key, col, row_offset=row_offset)
            )

        repeat_rows = 1 if include_header else 0
        table = Table(table_data, colWidths=COL_WIDTHS, repeatRows=repeat_rows)
        data_start = 1 if include_header else 0
        table.setStyle(
            TableStyle(
                [
                    *(
                        [("BACKGROUND", (0, 0), (-1, 0), HEADER_BG)]
                        if include_header
                        else []
                    ),
                    ("VALIGN", (0, data_start), (-1, -1), "TOP"),
                    ("VALIGN", (0, data_start), (2, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.25, LINE),
                    ("LINEAFTER", (0, 0), (-2, -1), 0.25, LINE),
                    *span_styles,
                ]
            )
        )
        return table

    groups = group_by_codigo(rows)
    header_used = False
    for group_rows in groups:
        for chunk in chunk_group_rows(group_rows):
            include_header = not header_used
            story.append(build_group_table(chunk, include_header=include_header))
            header_used = True
            story.append(Spacer(1, 1 * mm))

    story.append(Spacer(1, 2 * mm))

    def footer(canvas, doc_obj):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(MUTED)
        canvas.drawString(
            doc_obj.leftMargin,
            4 * mm,
            f"P38 · Catálogo 4×3{footer_tag if sem_estoque else ''} · {len(rows)} SKUs · A4 retrato",
        )
        canvas.drawRightString(
            PAGE_SIZE[0] - doc_obj.rightMargin,
            4 * mm,
            f"pág. {canvas.getPageNumber()}",
        )
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def main() -> int:
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUT
    filter_path = Path(sys.argv[3]) if len(sys.argv) > 3 else None

    if not xlsx.is_file():
        print(f"Ficheiro em falta: {xlsx}", file=sys.stderr)
        return 1

    codigos_filter = load_filter_codigos(filter_path)
    filter_kind = load_filter_kind(filter_path)
    sem_estoque = codigos_filter is not None
    rows = collapse_groups(load_catalog_rows(xlsx, codigos_filter=codigos_filter))
    build_pdf(rows, out, sem_estoque=sem_estoque, filter_kind=filter_kind)
    label = filter_kind or ("sem-estoque" if sem_estoque else "catalogo-4x3")
    print(f"[pdf:{label}] {len(rows)} SKUs → {out}")

    if not sem_estoque:
        try:
            ARTIFACT_OUT.parent.mkdir(parents=True, exist_ok=True)
            ARTIFACT_OUT.write_bytes(out.read_bytes())
            print(f"[pdf:catalogo-4x3] cópia → {ARTIFACT_OUT}")
        except OSError:
            pass

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
